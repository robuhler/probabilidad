import argparse
import csv
import unicodedata
from datetime import date, datetime
from pathlib import Path
from typing import Iterable


DATE_ALIASES = {
    "fecha",
    "fecha pizarra",
    "fecha de pizarra",
    "fecha de operacion",
    "fecha operacion",
}
PRICE_ALIASES = {
    "precio",
    "precio pizarra",
    "precio estimativo",
    "precio pizarra y estimativo",
    "pizarra y estimativo",
    "pizarra estimativo",
}
OUTPUT_COLUMNS = ("date", "price_ars_tn", "source", "price_kind", "product")
HEADER_SCAN_LIMIT = 50


def normalize_text(value: object) -> str:
    if value is None:
        return ""
    text = str(value).strip().lower()
    text = unicodedata.normalize("NFKD", text)
    return "".join(char for char in text if not unicodedata.combining(char))


def choose_excel_file() -> Path:
    try:
        import tkinter as tk
        from tkinter import filedialog
    except Exception as exc:  # pragma: no cover - depends on local GUI availability
        raise RuntimeError(
            "No pude abrir el selector de archivos. Ejecutá el script con --input 'ruta/al/archivo.xlsx'."
        ) from exc

    root = tk.Tk()
    root.withdraw()
    filename = filedialog.askopenfilename(
        title="Seleccioná el Excel de CAC/BCR para normalizar",
        filetypes=[("Excel", "*.xlsx *.xlsm *.xltx *.xltm")],
    )
    root.destroy()

    if not filename:
        raise RuntimeError("No se seleccionó ningún archivo Excel.")

    return Path(filename)


def detect_header(rows: Iterable[tuple[object, ...]]) -> tuple[int, dict[str, int]]:
    for index, row in enumerate(rows):
        normalized = [normalize_text(cell) for cell in row]
        mapping: dict[str, int] = {}
        for column_index, cell in enumerate(normalized):
            if cell in DATE_ALIASES:
                mapping["date"] = column_index
            if cell in PRICE_ALIASES:
                mapping["price_ars_tn"] = column_index
        if {"date", "price_ars_tn"}.issubset(mapping):
            return index, mapping
    raise ValueError(
        "No pude identificar automáticamente las columnas de fecha y precio en el Excel. "
        "Revisá los encabezados reales del archivo descargado."
    )


def parse_excel_date(value: object) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%Y-%m-%d %H:%M:%S", "%d/%m/%Y %H:%M:%S"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            continue
    raise ValueError(f"No pude interpretar la fecha: {value}")


def parse_price(value: object) -> str:
    if value is None:
        raise ValueError("El precio está vacío.")
    if isinstance(value, (int, float)):
        return f"{float(value):.2f}"

    text = str(value).strip().replace(" ", "")
    if not text:
        raise ValueError("El precio está vacío.")

    if "," in text:
        normalized = text.replace(".", "").replace(",", ".")
    else:
        normalized = text

    return f"{float(normalized):.2f}"


def normalize_rows(rows: list[tuple[object, ...]], source: str = "CAC BCR") -> list[dict[str, str]]:
    header_index, mapping = detect_header(rows[:HEADER_SCAN_LIMIT])
    normalized_rows: list[dict[str, str]] = []
    for row in rows[header_index + 1 :]:
        if row is None:
            continue
        date_value = row[mapping["date"]] if mapping["date"] < len(row) else None
        price_value = row[mapping["price_ars_tn"]] if mapping["price_ars_tn"] < len(row) else None
        if date_value in (None, "") and price_value in (None, ""):
            continue
        normalized_rows.append(
            {
                "date": parse_excel_date(date_value),
                "price_ars_tn": parse_price(price_value),
                "source": source,
                "price_kind": "Precio Pizarra y Estimativo",
                "product": "Trigo",
            }
        )
    return normalized_rows


def main() -> None:
    parser = argparse.ArgumentParser(description="Normaliza un Excel descargado desde CAC/BCR a CSV estándar.")
    parser.add_argument("--input", type=Path, required=False, help="Archivo Excel descargado desde CAC/BCR")
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("trigo_prediccion/spot_trigo/normalized/precio_spot_trigo_normalizado.csv"),
        help="CSV normalizado de salida",
    )
    parser.add_argument("--sheet", type=str, default=None, help="Nombre de hoja, si querés forzar una hoja")
    args = parser.parse_args()

    input_path = args.input if args.input else choose_excel_file()

    from openpyxl import load_workbook

    workbook = load_workbook(input_path, data_only=True)
    sheet = workbook[args.sheet] if args.sheet else workbook[workbook.sheetnames[0]]
    rows = list(sheet.iter_rows(values_only=True))
    normalized_rows = normalize_rows(rows)

    args.output.parent.mkdir(parents=True, exist_ok=True)
    with args.output.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(OUTPUT_COLUMNS))
        writer.writeheader()
        writer.writerows(normalized_rows)

    print(f"Archivo normalizado generado: {args.output} ({len(normalized_rows)} filas)")


if __name__ == "__main__":
    main()
