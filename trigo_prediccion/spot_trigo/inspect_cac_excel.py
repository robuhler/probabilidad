import argparse
from pathlib import Path
from typing import Iterable



def compact_row(row: Iterable[object]) -> list[str]:
    values = []
    for cell in row:
        text = "" if cell is None else str(cell).strip()
        values.append(text)
    return values


def choose_excel_file() -> Path:
    try:
        import tkinter as tk
        from tkinter import filedialog
    except Exception as exc:  # pragma: no cover - depends on local GUI availability
        raise RuntimeError(
            "No pude abrir el selector de archivos. Ejecutá el script con --input 'ruta/al/archivo.xlsx'."
        ) from exc

    root = tk.Tk()
    root.withdraw()
    filename = filedialog.askopenfilename(
        title="Seleccioná el Excel de CAC/BCR",
        filetypes=[("Excel", "*.xlsx *.xlsm *.xltx *.xltm")],
    )
    root.destroy()

    if not filename:
        raise RuntimeError("No se seleccionó ningún archivo Excel.")

    return Path(filename)


def main() -> None:
    parser = argparse.ArgumentParser(description="Inspecciona un Excel de CAC/BCR para entender su estructura.")
    parser.add_argument("--input", type=Path, required=False, help="Ruta al archivo Excel")
    parser.add_argument("--rows", type=int, default=12, help="Cantidad de filas a mostrar por hoja")
    args = parser.parse_args()

    input_path = args.input if args.input else choose_excel_file()

    from openpyxl import load_workbook

    workbook = load_workbook(input_path, data_only=True)
    print(f"Archivo: {input_path}")
    print(f"Hojas detectadas: {', '.join(workbook.sheetnames)}")

    for sheet_name in workbook.sheetnames:
        sheet = workbook[sheet_name]
        print("\n" + "=" * 80)
        print(f"Hoja: {sheet_name}")
        print("=" * 80)
        shown = 0
        for row_index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            compact = compact_row(row)
            if not any(compact):
                continue
            print(f"Fila {row_index}: {compact}")
            shown += 1
            if shown >= args.rows:
                break


if __name__ == "__main__":
    main()
